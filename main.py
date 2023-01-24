from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image

from classes.models import LeitorDeAcoes, GerenciadorPlanilha

acao = input('Qual o código da ação a ser processada? ').upper()

# Busca o arquivo do Código da Ação informado pelo usuário
leitor_acoes = LeitorDeAcoes(caminho_arquivo='./assets/dados/')
leitor_acoes.processa_arquivo(acao)

# Cria um Objeto gerenciador de planilhas
gerenciador = GerenciadorPlanilha()
planilha_dados = gerenciador.adicionar_planilha("Dados")

# PREPARANDO A PLANILHA DE 'DADOS'
gerenciador.adiciona_linha(["DATA", "COTAÇÃO", "BANDA INFERIOR", "BANDA SUPERIOR"])

indice = 2
for linha in leitor_acoes.dados:
    # TRATANDO E PREPARANDO A DATA, CRIANDO UM OBJETO DATE DO PYTHON
    ano_mes_dia = linha[0].split(" ")[0]
    data = date(
        year=int(ano_mes_dia.split("-")[0]),
        month=int(ano_mes_dia.split("-")[1]),
        day=int(ano_mes_dia.split("-")[2])
    )

    # TRATANDO E PREPARANDO A COTAÇÃO E AS FÓRMULAS DA BANDA DE BOLLINGER
    cotacao = float(linha[1])
    formula_bb_inferior = f'=AVERAGE(B{indice}:B{indice + 19}) - 2*STDEV(B{indice}:B{indice + 19})'
    formula_bb_superior = f'=AVERAGE(B{indice}:B{indice + 19}) + 2*STDEV(B{indice}:B{indice + 19})'

    # ESCREVENDO DADOS NA PLANILHA ATIVA
    # Data
    gerenciador.atualiza_celula(celula=f'A{indice}', dado=data)
    # Cotação
    gerenciador.atualiza_celula(celula=f'B{indice}', dado=cotacao)
    # Banda inferior de Bollinger
    gerenciador.atualiza_celula(celula=f'C{indice}', dado=formula_bb_inferior)
    # Banda superior de Bollinger
    gerenciador.atualiza_celula(celula=f'D{indice}', dado=formula_bb_superior)

    # PREPARA PARA A ESCRITA NA PRÓXIMA LINHA
    indice += 1

planilha_grafico = gerenciador.adicionar_planilha(titulo_planilha="Gráfico")

# Mesclagem de células e criação de estilos, preenchimento e definição de alinhamento
gerenciador.mescla_celulas(celula_inicio="A1", celula_fim="T2")
gerenciador.estiliza_fonte(celula="A1", fonte=Font(b=True, sz=18, color="FFFFFF"))
gerenciador.estiliza_preenchimento(celula="A1", preenchimento=PatternFill("solid", fgColor="07838f"))
gerenciador.estiliza_alinhamento(celula="A1", alinhamento=Alignment(vertical="center", horizontal="center"))
gerenciador.atualiza_celula(celula="A1", dado="Histórico de Cotações")

# Criação do Gráfico
grafico = LineChart()
grafico.width = 33.95
grafico.height = 15.5
grafico.title = f"Cotações - {acao}"
grafico.x_axis.title = "Data da Cotação"
grafico.y_axis.title = "Valor da Cotação"

# Criando referencias e adicionando elas como séries de dados e datas
referencia_cotacoes = Reference(planilha_ativa, min_col=2, min_row=2, max_col=4, max_row=indice)
referencia_datas = Reference(planilha_ativa, min_col=1, min_row=2, max_col=1, max_row=indice)
grafico.add_data(referencia_cotacoes)
grafico.set_categories(referencia_datas)

# Editando as linhas dos gráficos, os estilos e cores
linha_cotacao = grafico.series[0]
linha_bb_inferior = grafico.series[1]
linha_bb_superior = grafico.series[2]

linha_cotacao.graphicalProperties.line.width = 0
linha_cotacao.graphicalProperties.line.solidFill = "000000"

linha_bb_inferior.graphicalProperties.line.width = 0
linha_bb_inferior.graphicalProperties.line.solidFill = "09ab11"

linha_bb_superior.graphicalProperties.line.width = 0
linha_bb_superior.graphicalProperties.line.solidFill = "ff0000"

# Adicionando de fato o gráfico à planilha
planilha_grafico.add_chart(grafico, "A3")

# Adicionando uma imagem
imagem = Image("./assets/recursos/logo.png")
planilha_grafico.merge_cells("I33:L36")
planilha_grafico.add_image(imagem, "I33")

# Gerando o arquivo XLSX final
workbook.save(f'./saida/{acao}.xlsx')
